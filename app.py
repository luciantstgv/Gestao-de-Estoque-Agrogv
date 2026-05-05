from datetime import date, datetime
import csv
import io
import sqlite3
from functools import wraps
from flask import Flask, g, redirect, render_template, request, session, url_for
from openpyxl import load_workbook

app = Flask(__name__)
app.secret_key = "agrogv-mvp-secret"
DB_PATH = "estoque.db"

ROLE_PERMS = {
    "Administrador": {"dashboard", "products", "lots", "import", "sales", "movements", "alerts", "reports"},
    "Operador": {"dashboard", "lots", "sales", "movements", "alerts", "products"},
    "Gestor": {"dashboard", "alerts", "reports", "products"},
}

IMPORT_FIELDS = ["codigo_interno","codigo_barras","nome_produto","descricao","categoria","unidade_medida","marca_fabricante","fornecedor","valor_custo","valor_venda","estoque_minimo","numero_lote","quantidade_inicial_lote","quantidade_atual_lote","data_fabricacao","data_validade","localizacao_fisica","responsavel_conferencia","data_conferencia","observacoes"]


def now_iso(): return datetime.utcnow().isoformat()

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db

@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None: db.close()


def init_db():
    db = get_db()
    db.executescript("""
    CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT,username TEXT UNIQUE,password TEXT NOT NULL,role TEXT NOT NULL,status TEXT NOT NULL DEFAULT 'active');
    CREATE TABLE IF NOT EXISTS products (id INTEGER PRIMARY KEY AUTOINCREMENT,internal_code TEXT UNIQUE NOT NULL,barcode TEXT,name TEXT NOT NULL,description TEXT,category TEXT,unit TEXT,brand TEXT,supplier TEXT,cost_value REAL DEFAULT 0,sale_value REAL DEFAULT 0,min_stock REAL DEFAULT 0,notes TEXT,status TEXT NOT NULL DEFAULT 'active',created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS lots (id INTEGER PRIMARY KEY AUTOINCREMENT,product_id INTEGER NOT NULL,lot_number TEXT NOT NULL,quantity_initial REAL NOT NULL,quantity_current REAL NOT NULL CHECK (quantity_current >= 0),manufacturing_date TEXT,expiry_date TEXT NOT NULL,location TEXT NOT NULL,checked_by TEXT,checked_at TEXT,notes TEXT,status TEXT NOT NULL DEFAULT 'active',created_at TEXT NOT NULL,FOREIGN KEY(product_id) REFERENCES products(id));
    CREATE TABLE IF NOT EXISTS sales (id INTEGER PRIMARY KEY AUTOINCREMENT,sale_date TEXT NOT NULL,sale_number TEXT NOT NULL,client_name TEXT,product_id INTEGER NOT NULL,quantity REAL NOT NULL,unit_value REAL NOT NULL,total_value REAL NOT NULL,seller TEXT,notes TEXT,created_at TEXT NOT NULL,FOREIGN KEY(product_id) REFERENCES products(id));
    CREATE TABLE IF NOT EXISTS movements (id INTEGER PRIMARY KEY AUTOINCREMENT,moved_at TEXT NOT NULL,product_id INTEGER,lot_id INTEGER,quantity REAL NOT NULL,movement_type TEXT NOT NULL,user_name TEXT,previous_balance REAL NOT NULL,posterior_balance REAL NOT NULL,observation TEXT,reference TEXT,FOREIGN KEY(product_id) REFERENCES products(id),FOREIGN KEY(lot_id) REFERENCES lots(id));
    """)
    defaults = [("admin", "admin123", "Administrador"), ("operador", "operador123", "Operador"), ("gestor", "gestor123", "Gestor")]
    for u,p,r in defaults:
        db.execute("INSERT OR IGNORE INTO users(username,password,role,status) VALUES(?,?,?,'active')", (u,p,r))
    db.commit()


def to_float(v, d=0):
    try: return float(v)
    except: return d

def validade_status(expiry_date):
    days = (date.fromisoformat(expiry_date) - date.today()).days
    if days < 0: return "Vencido", days
    if days <= 60: return "Crítico", days
    if days <= 90: return "Atenção", days
    return "Normal", days


def create_movement(product_id, lot_id, quantity, movement_type, user_name, previous_balance, posterior_balance, observation, reference):
    get_db().execute("INSERT INTO movements(moved_at,product_id,lot_id,quantity,movement_type,user_name,previous_balance,posterior_balance,observation,reference) VALUES(?,?,?,?,?,?,?,?,?,?)", (now_iso(), product_id, lot_id, quantity, movement_type, user_name, previous_balance, posterior_balance, observation, reference))


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("user"): return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapped

def role_required(page):
    def decorator(view):
        @wraps(view)
        def wrapped(*args, **kwargs):
            user = session.get("user")
            if not user: return redirect(url_for("login"))
            if page not in ROLE_PERMS.get(user["role"], set()): return "Acesso negado", 403
            return view(*args, **kwargs)
        return wrapped
    return decorator

@app.context_processor
def inject_user(): return {"logged_user": session.get("user")}

@app.route('/login', methods=['GET','POST'])
def login():
    error = None
    if request.method == 'POST':
        user = get_db().execute("SELECT * FROM users WHERE username=? AND password=? AND status='active'", (request.form.get("username"), request.form.get("password"))).fetchone()
        if user: 
            session["user"] = {"username": user["username"], "role": user["role"]}
            return redirect(url_for('dashboard'))
        error = "Credenciais inválidas"
    return render_template("login.html", error=error)

@app.route('/logout')
def logout():
    session.clear(); return redirect(url_for('login'))

@app.route('/')
@login_required
@role_required("dashboard")
def dashboard():
    db = get_db()
    total_products = db.execute("SELECT COUNT(*) c FROM products WHERE status='active'").fetchone()["c"]
    total_lots = db.execute("SELECT COUNT(*) c FROM lots WHERE status='active'").fetchone()["c"]
    lot_rows = db.execute("SELECT l.*, p.cost_value FROM lots l JOIN products p ON p.id=l.product_id WHERE l.status='active'").fetchall()
    expired=v30=v60=v90=0; risk=0
    for l in lot_rows:
        _, days = validade_status(l["expiry_date"])
        if days < 0: expired += 1
        if 0 <= days <= 30: v30 += 1; risk += l["quantity_current"] * l["cost_value"]
        if 0 <= days <= 60: v60 += 1
        if 0 <= days <= 90: v90 += 1
    low_stock = db.execute("SELECT COUNT(*) c FROM (SELECT p.id,COALESCE(SUM(l.quantity_current),0) q,p.min_stock FROM products p LEFT JOIN lots l ON l.product_id=p.id AND l.status='active' WHERE p.status='active' GROUP BY p.id HAVING q<=p.min_stock)").fetchone()["c"]
    return render_template('dashboard.html', stats=locals())

@app.route('/products', methods=['GET','POST'])
@login_required
@role_required("products")
def products_page():
    db = get_db()
    if request.method == 'POST' and session["user"]["role"] in ["Administrador", "Operador"]:
        db.execute("INSERT INTO products(internal_code,barcode,name,description,category,unit,brand,supplier,cost_value,sale_value,min_stock,notes,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                   (request.form.get("internal_code"), request.form.get("barcode"), request.form.get("name"), request.form.get("description"), request.form.get("category"), request.form.get("unit"), request.form.get("brand"), request.form.get("supplier"), to_float(request.form.get("cost_value")), to_float(request.form.get("sale_value")), to_float(request.form.get("min_stock")), request.form.get("notes"), "active", now_iso()))
        db.commit(); return redirect(url_for('products_page'))
    return render_template('products.html', products=db.execute("SELECT * FROM products ORDER BY id DESC").fetchall())

@app.route('/lots', methods=['GET','POST'])
@login_required
@role_required("lots")
def lots_page():
    db=get_db()
    if request.method=='POST':
        qi, qc = to_float(request.form.get("quantity_initial")), to_float(request.form.get("quantity_current"), to_float(request.form.get("quantity_initial")))
        db.execute("INSERT INTO lots(product_id,lot_number,quantity_initial,quantity_current,manufacturing_date,expiry_date,location,checked_by,checked_at,notes,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                   (request.form.get("product_id"), request.form.get("lot_number"), qi, qc, request.form.get("manufacturing_date") or None, request.form.get("expiry_date"), request.form.get("location"), request.form.get("checked_by"), request.form.get("checked_at") or None, request.form.get("notes"), "active", now_iso()))
        lot_id = db.execute("SELECT last_insert_rowid() id").fetchone()["id"]
        create_movement(int(request.form.get("product_id")), lot_id, qc, "entrada", session["user"]["username"], 0, qc, "Cadastro de lote", "lot_create")
        db.commit(); return redirect(url_for('lots_page'))
    lots = db.execute("SELECT l.*,p.name product_name FROM lots l JOIN products p ON p.id=l.product_id ORDER BY l.expiry_date").fetchall()
    rows=[dict(r)|{"validity":validade_status(r["expiry_date"])[0]} for r in lots]
    return render_template('lots.html', lots=rows, products=db.execute("SELECT id,name,internal_code FROM products WHERE status='active' ORDER BY name").fetchall())

@app.route('/sales', methods=['GET','POST'])
@login_required
@role_required("sales")
def sales_page():
    db=get_db(); error=None
    if request.method=='POST':
        pid=int(request.form.get("product_id")); qty=to_float(request.form.get("quantity")); uv=to_float(request.form.get("unit_value"))
        lots=db.execute("SELECT * FROM lots WHERE product_id=? AND status='active' AND quantity_current>0 ORDER BY expiry_date",(pid,)).fetchall(); avail=sum(l["quantity_current"] for l in lots)
        if qty>avail: error=f"Estoque insuficiente. Disponível: {avail:.2f}."
        else:
            total=qty*uv
            db.execute("INSERT INTO sales(sale_date,sale_number,client_name,product_id,quantity,unit_value,total_value,seller,notes,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                       (request.form.get("sale_date"), request.form.get("sale_number"), request.form.get("client_name"), pid, qty, uv, total, request.form.get("seller"), request.form.get("notes"), now_iso()))
            sale_id=db.execute("SELECT last_insert_rowid() id").fetchone()["id"]; rem=qty
            for l in lots:
                if rem<=0: break
                take=min(rem,l["quantity_current"]); prev=l["quantity_current"]; post=prev-take
                db.execute("UPDATE lots SET quantity_current=? WHERE id=?",(post,l["id"])); create_movement(pid,l["id"],-take,"venda",session["user"]["username"],prev,post,"Baixa PVPS",f"venda:{sale_id}"); rem-=take
            db.commit(); return redirect(url_for('sales_page'))
    products=db.execute("SELECT id,name,internal_code FROM products WHERE status='active' ORDER BY name").fetchall(); sales=db.execute("SELECT s.*,p.name product_name FROM sales s JOIN products p ON p.id=s.product_id ORDER BY s.id DESC").fetchall()
    return render_template('sales.html', products=products, sales=sales, error=error)


def parse_upload(fs):
    if fs.filename.lower().endswith('.csv'):
        return list(csv.DictReader(io.StringIO(fs.stream.read().decode('utf-8-sig'))))
    wb=load_workbook(fs,data_only=True); ws=wb.active; h=[str(c.value).strip() if c.value else '' for c in ws[1]]
    return [{h[i]: row[i] for i in range(len(h))} for row in ws.iter_rows(min_row=2,values_only=True)]

@app.route('/import', methods=['GET','POST'])
@login_required
@role_required("import")
def import_page():
    preview=session.get('import_preview',[]); errors=session.get('import_errors',[]); report=session.get('import_report')
    if request.method=='POST':
        if request.form.get('action')=='preview':
            cleaned=[]; errs=[]
            for i,r in enumerate(parse_upload(request.files['file']),start=2):
                m={k:str(r.get(k,'') or '').strip() for k in IMPORT_FIELDS}
                if not m['codigo_interno'] or not m['nome_produto'] or not m['numero_lote'] or not m['data_validade']: errs.append(f"Linha {i}: campos obrigatórios ausentes")
                else: cleaned.append(m)
            session['import_preview']=cleaned[:20]; session['import_full']=cleaned; session['import_errors']=errs; return redirect(url_for('import_page'))
        if request.form.get('action')=='confirm':
            db=get_db(); cp=up=cl=0; errs=[]
            for r in session.get('import_full',[]):
                try:
                    p=db.execute("SELECT id FROM products WHERE internal_code=?",(r['codigo_interno'],)).fetchone()
                    if p:
                        db.execute("UPDATE products SET barcode=?,name=?,description=?,category=?,unit=?,brand=?,supplier=?,cost_value=?,sale_value=?,min_stock=?,notes=? WHERE id=?",(r['codigo_barras'],r['nome_produto'],r['descricao'],r['categoria'],r['unidade_medida'],r['marca_fabricante'],r['fornecedor'],to_float(r['valor_custo']),to_float(r['valor_venda']),to_float(r['estoque_minimo']),r['observacoes'],p['id'])); pid=p['id']; up+=1
                    else:
                        db.execute("INSERT INTO products(internal_code,barcode,name,description,category,unit,brand,supplier,cost_value,sale_value,min_stock,notes,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",(r['codigo_interno'],r['codigo_barras'],r['nome_produto'],r['descricao'],r['categoria'],r['unidade_medida'],r['marca_fabricante'],r['fornecedor'],to_float(r['valor_custo']),to_float(r['valor_venda']),to_float(r['estoque_minimo']),r['observacoes'],'active',now_iso())); pid=db.execute("SELECT last_insert_rowid() id").fetchone()['id']; cp+=1
                    qi,qc=to_float(r['quantidade_inicial_lote']),to_float(r['quantidade_atual_lote'])
                    db.execute("INSERT INTO lots(product_id,lot_number,quantity_initial,quantity_current,manufacturing_date,expiry_date,location,checked_by,checked_at,notes,status,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",(pid,r['numero_lote'],qi,qc if qc>0 else qi,r['data_fabricacao'] or None,r['data_validade'],r['localizacao_fisica'] or 'N/A',r['responsavel_conferencia'],r['data_conferencia'] or None,r['observacoes'],'active',now_iso())); lid=db.execute("SELECT last_insert_rowid() id").fetchone()['id']
                    amount=qc if qc>0 else qi; create_movement(pid,lid,amount,'entrada',session['user']['username'],0,amount,'Importação','import'); cl+=1
                except Exception as exc: errs.append(str(exc))
            db.commit(); session['import_report']={"created_products":cp,"updated_products":up,"created_lots":cl,"errors":errs}; session.pop('import_full',None); return redirect(url_for('import_page'))
    return render_template('import.html', preview=preview, errors=errors, report=report)

@app.route('/movements')
@login_required
@role_required("movements")
def movements_page():
    rows=get_db().execute("SELECT m.*,p.name product_name,l.lot_number FROM movements m LEFT JOIN products p ON p.id=m.product_id LEFT JOIN lots l ON l.id=m.lot_id ORDER BY m.id DESC LIMIT 500").fetchall()
    return render_template('movements.html', movements=rows)

@app.route('/alerts')
@login_required
@role_required("alerts")
def alerts_page():
    rows=get_db().execute("SELECT l.*,p.name product_name FROM lots l JOIN products p ON p.id=l.product_id WHERE l.status='active' ORDER BY l.expiry_date").fetchall()
    alerts=[dict(r)|{"validity":validade_status(r['expiry_date'])[0],"days":validade_status(r['expiry_date'])[1]} for r in rows if validade_status(r['expiry_date'])[0] != 'Normal']
    return render_template('alerts.html', alerts=alerts)

@app.route('/reports')
@login_required
@role_required("reports")
def reports_page():
    db=get_db()
    expired_30_60_90 = db.execute("SELECT l.*,p.name product_name,p.cost_value FROM lots l JOIN products p ON p.id=l.product_id WHERE l.status='active' ORDER BY l.expiry_date").fetchall()
    expired=[];r30=[];r60=[];r90=[];risk=0
    for r in expired_30_60_90:
        _,d=validade_status(r['expiry_date'])
        data=dict(r)|{"days":d}
        if d<0: expired.append(data)
        if 0<=d<=30: r30.append(data); risk += r['quantity_current']*r['cost_value']
        if 0<=d<=60: r60.append(data)
        if 0<=d<=90: r90.append(data)
    low_stock=db.execute("SELECT p.internal_code,p.name,COALESCE(SUM(l.quantity_current),0) stock,p.min_stock FROM products p LEFT JOIN lots l ON l.product_id=p.id AND l.status='active' WHERE p.status='active' GROUP BY p.id HAVING stock<=p.min_stock").fetchall()
    no_lot=db.execute("SELECT p.* FROM products p LEFT JOIN lots l ON l.product_id=p.id AND l.status='active' WHERE p.status='active' GROUP BY p.id HAVING COUNT(l.id)=0").fetchall()
    sales=db.execute("SELECT s.*,p.name product_name FROM sales s JOIN products p ON p.id=s.product_id ORDER BY s.sale_date DESC,s.id DESC").fetchall()
    movements=db.execute("SELECT m.*,p.name product_name,l.lot_number FROM movements m LEFT JOIN products p ON p.id=m.product_id LEFT JOIN lots l ON l.id=m.lot_id ORDER BY m.id DESC LIMIT 300").fetchall()
    return render_template('reports.html', expired=expired, r30=r30, r60=r60, r90=r90, low_stock=low_stock, no_lot=no_lot, sales=sales, movements=movements, risk=risk)

if __name__ == '__main__':
    with app.app_context(): init_db()
    app.run(debug=True)
